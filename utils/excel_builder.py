import pandas as pd
from io import BytesIO
import tempfile
import shutil
from openpyxl import load_workbook
import datetime as dt


def build_excel(template_path: str, 갑지_데이터: dict, 시트_데이터: dict) -> BytesIO:
    """
    범용 Excel 생성 함수
    
    Args:
        template_path: 템플릿 파일 경로
        갑지_데이터: {"용역명": "...", "발주기관명": "...", "도급예정액": 123}
        시트_데이터: {"내역서": df, "투입인원 및 내역": df, ...}
    
    Returns:
        BytesIO: 생성된 Excel 파일
    """
    # 템플릿 복사
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    shutil.copy(template_path, tmp.name)
    
    # 갑지 시트 채우기
    wb = load_workbook(tmp.name)
    ws = wb["갑지"]
    ws["A1"].value = dt.date.today().strftime("%Y-%m-%d")
    ws["D10"].value = 갑지_데이터.get("용역명", "")
    ws["G22"].value = 갑지_데이터.get("발주기관명", "")
    raw = 갑지_데이터.get("도급예정액", 0)
    ws["G20"].value = f"{int(raw // 1000) * 1000:,} 원"
    wb.save(tmp.name)
    
    # 나머지 시트 덮어쓰기
    with pd.ExcelWriter(tmp.name, engine="openpyxl", mode="a", if_sheet_exists="overlay") as writer:
        for sheet_name, df in 시트_데이터.items():
            if df is not None and not df.empty:
                df.to_excel(writer, sheet_name=sheet_name, index=False, header=False, startrow=2)
    
    # BytesIO 반환
    buf = BytesIO()
    with open(tmp.name, "rb") as f:
        buf.write(f.read())
    buf.seek(0)
    return buf
